import os
import re
import openpyxl
import xlrd
from openpyxl.worksheet.page import PageMargins
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox
from tkinter import ttk
import threading

CM_TO_POINTS = 28.35


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("工资表打印格式批量调整工具")
        self.root.geometry("700x500")
        self.file_list = []

        # 文件列表区域
        tk.Label(root, text="待处理文件列表：", font=("微软雅黑", 10)).pack(anchor='w', padx=10, pady=5)
        self.listbox = tk.Listbox(root, selectmode=tk.EXTENDED, height=8)
        self.listbox.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        # 按钮区域
        btn_frame = tk.Frame(root)
        btn_frame.pack(fill=tk.X, padx=10, pady=5)
        tk.Button(btn_frame, text="添加文件", command=self.add_files).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="清空列表", command=self.clear_list).pack(side=tk.LEFT, padx=5)
        tk.Button(btn_frame, text="开始处理", command=self.start_process).pack(side=tk.RIGHT, padx=5)

        # 进度条
        self.progress = ttk.Progressbar(root, orient='horizontal', length=400, mode='determinate')
        self.progress.pack(fill=tk.X, padx=10, pady=5)

        # 日志区域
        tk.Label(root, text="处理日志：", font=("微软雅黑", 10)).pack(anchor='w', padx=10)
        self.log = scrolledtext.ScrolledText(root, height=12, state='normal')
        self.log.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

    def add_files(self):
        files = filedialog.askopenfilenames(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls")]
        )
        for f in files:
            if f not in self.file_list:
                self.file_list.append(f)
                self.listbox.insert(tk.END, f)

    def clear_list(self):
        self.file_list.clear()
        self.listbox.delete(0, tk.END)

    def start_process(self):
        if not self.file_list:
            messagebox.showwarning("提示", "请先添加文件！")
            return
        self.disable_buttons()
        self.progress['value'] = 0
        self.log.delete(1.0, tk.END)
        threading.Thread(target=self.process_files, daemon=True).start()

    def disable_buttons(self):
        for child in self.root.winfo_children():
            if isinstance(child, tk.Frame):
                for btn in child.winfo_children():
                    if isinstance(btn, tk.Button):
                        btn.config(state='disabled')

    def enable_buttons(self):
        for child in self.root.winfo_children():
            if isinstance(child, tk.Frame):
                for btn in child.winfo_children():
                    if isinstance(btn, tk.Button):
                        btn.config(state='normal')

    def process_files(self):
        total = len(self.file_list)
        for idx, file_path in enumerate(self.file_list):
            self.log_insert(f"正在处理: {os.path.basename(file_path)}")
            try:
                self.process_one_file(file_path)
                self.log_insert(f"处理成功: {os.path.basename(file_path)}")
            except Exception as e:
                self.log_insert(f"处理失败: {os.path.basename(file_path)} - {str(e)}")
            progress = int((idx + 1) / total * 100)
            self.root.after(0, lambda p=progress: self.progress.configure(value=p))
        self.root.after(0, lambda: messagebox.showinfo("完成", "全部处理完成！"))
        self.root.after(0, self.enable_buttons)

    def log_insert(self, msg):
        self.root.after(0, lambda: self.log.insert(tk.END, msg + "\n"))
        self.root.after(0, lambda: self.log.see(tk.END))

    def read_xls_file(self, file_path):
        """读取 .xls 文件并返回数据"""
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)

        # 创建新的 openpyxl 工作簿
        wb = openpyxl.Workbook()
        ws = wb.active

        # 复制数据
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                cell_value = sheet.cell_value(row, col)
                ws.cell(row + 1, col + 1, cell_value)

        return wb, ws

    def process_one_file(self, file_path):
        """处理 Excel 文件（支持 .xls 和 .xlsx）"""

        # 根据文件格式选择读取方式
        if file_path.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file_path, data_only=True)
            ws = wb.active
        elif file_path.endswith('.xls'):
            wb, ws = self.read_xls_file(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {file_path}")

        # 获取新文件名
        # 取第二行 B 列内容
        row2_colB = ws['B2'].value
        if row2_colB:
            name_part = str(row2_colB).strip()
            name_part = name_part.replace("名称：", "").strip()
        else:
            name_part = "未知名称"

        # 取第一行 C 列内容中的年月
        row1_colC = ws['C1'].value
        year_month = ""
        if row1_colC:
            match = re.search(r'(\d{4})年\s*(\d{1,2})\s*月', str(row1_colC))
            if match:
                year_month = f"{match.group(1)}年{match.group(2)}月"

        # 构建新文件名（统一保存为 .xlsx 格式）
        new_filename = f"{name_part}{year_month}工资表.xlsx"
        original_dir = os.path.dirname(file_path)
        new_file_path = os.path.join(original_dir, new_filename)

        # 处理重复文件名
        if os.path.exists(new_file_path):
            base, ext = os.path.splitext(new_file_path)
            counter = 1
            while os.path.exists(f"{base}_{counter}{ext}"):
                counter += 1
            new_file_path = f"{base}_{counter}{ext}"
            self.log_insert(f"文件已存在，另存为：{os.path.basename(new_file_path)}")

        self.log_insert(f"准备另存为：{os.path.basename(new_file_path)}")

        # 1. 设置页边距（2cm 左侧边距）
        ws.page_margins = PageMargins(
            left=2 / 2.54,  # 2cm 转英寸
            right=0.75,  # 默认右边距
            top=0.75,  # 默认上边距
            bottom=0.75,  # 默认下边距
            header=0.5,  # 页眉边距
            footer=0.5  # 页脚边距
        )

        # 2. 设置打印标题行
        ws.print_title_rows = '1:5'

        # 3. 设置缩放：宽度缩放到1页
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = False

        # 4. 隐藏D列和E列
        ws.column_dimensions['D'].hidden = True
        ws.column_dimensions['E'].hidden = True

        # 5. 修改签字文本
        last_row = ws.max_row
        last_col = ws.max_column

        found = False
        start_row = max(1, last_row - 2)
        for row in range(start_row, last_row + 1):
            for col in range(1, last_col + 1):
                cell = ws.cell(row, col)
                if cell.value and isinstance(cell.value, str):
                    if "部长、分管副总签字：" in cell.value:
                        new_val = cell.value.replace("部长、分管副总签字：", "部长签字：")
                        cell.value = new_val
                        self.log_insert(f"已修改签字文本: {new_val}")
                        found = True
                        break
            if found:
                break

        if not found:
            self.log_insert("警告：未找到‘部长、分管副总签字：’文本")

        # 6. 检查四列基数
        header_row = 3
        keywords = ["社保基数", "医保基数", "工伤基数", "公积金基数"]
        col_indices = []

        for col in range(1, min(51, ws.max_column + 1)):
            cell_val = ws.cell(header_row, col).value
            if cell_val and isinstance(cell_val, str):
                for kw in keywords:
                    if kw in cell_val:
                        col_indices.append(col)
                        break

        if len(col_indices) != 4:
            self.log_insert(f"警告：未找到全部四列基数，找到 {len(col_indices)} 列")
        else:
            all_zero = True
            for row in range(header_row + 1, last_row + 1):
                for col in col_indices:
                    val = ws.cell(row, col).value
                    if val is not None:
                        try:
                            if float(val) != 0:
                                all_zero = False
                                break
                        except (ValueError, TypeError):
                            all_zero = False
                            break
                if not all_zero:
                    break

            if all_zero:
                removed = False
                for row in range(last_row, max(1, last_row - 2), -1):
                    for col in range(1, last_col + 1):
                        cell = ws.cell(row, col)
                        if cell.value and "业务审核：" in str(cell.value):
                            new_val = str(cell.value).replace("业务审核：", "")
                            cell.value = new_val
                            removed = True
                            self.log_insert("已删除‘业务审核：’字样（条件满足）")
                            break
                    if removed:
                        break
                if not removed:
                    self.log_insert("警告：未找到‘业务审核：’文本")
            else:
                self.log_insert("四列基数不全为0，保留‘业务审核：’字样")

        # 保存为新文件
        wb.save(new_file_path)
        self.log_insert(f"处理完成，已保存为：{os.path.basename(new_file_path)}")


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()