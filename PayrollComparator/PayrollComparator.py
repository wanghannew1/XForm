#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
工资表人员差异对比工具
支持命令行和GUI两种模式
"""

import argparse
import json
import os
import sys
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

SCRIPT_DIR = Path(__file__).parent
CONFIG_FILE = SCRIPT_DIR / "config.json"
DEFAULT_TEMPLATE_FILE = SCRIPT_DIR / "员工导入模板202512.xlsx"


def load_config():
    """加载配置文件"""
    if CONFIG_FILE.exists():
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_config(config):
    """保存配置文件"""
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def read_salary_table(filepath, logger=None):
    """读取工资表，提取人员信息"""
    def log(msg):
        if logger:
            logger(msg)

    log(f"[INFO] 读取工资表: {filepath}")
    df = pd.read_excel(filepath, header=None)
    
    data = df.iloc[5:, [1, 2]].copy()
    data.columns = ["姓名", "身份证"]
    data = data.dropna(subset=["身份证"])
    data["身份证"] = data["身份证"].astype(str).str.strip()
    data = data[data["身份证"].str.match(r'^\d{17}[\dXx]$|^\d{15}$', na=False)]
    
    log(f"[INFO] 成功读取 {len(data)} 条人员记录")
    return data


def read_phone_table(filepath, logger=None):
    """读取电话表，返回身份证->手机号字典"""
    def log(msg):
        if logger:
            logger(msg)
    
    log(f"[INFO] 读取电话信息表: {filepath}")
    df = pd.read_excel(filepath)
    df.columns = df.columns.str.strip()
    
    id_col = None
    phone_col = None
    for col in df.columns:
        if "证件号码" in col or "身份证" in col:
            id_col = col
        if "手机" in col or "电话" in col:
            phone_col = col
    
    if not id_col or not phone_col:
        log("[ERROR] 电话表缺少必要列（证件号码、手机号）")
        return {}
    
    phone_dict = {}
    for _, row in df.iterrows():
        id_num = str(row[id_col]).strip()
        phone = str(row[phone_col]).strip()
        if phone and '.' in phone and phone.endswith('.0'):
            phone = phone[:-2]
        if id_num and id_num != "nan":
            phone_dict[id_num] = phone
    
    log(f"[INFO] 成功读取 {len(phone_dict)} 条电话记录")
    return phone_dict


def compare_persons(prev_df, curr_df, logger=None):
    """对比人员差异"""
    def log(msg):
        if logger:
            logger(msg)
    
    prev_ids = set(prev_df["身份证"])
    curr_ids = set(curr_df["身份证"])
    
    new_ids = curr_ids - prev_ids
    left_ids = prev_ids - curr_ids
    
    log(f"[INFO] 新增人员: {len(new_ids)} 人")
    log(f"[INFO] 减少人员: {len(left_ids)} 人")
    
    return {
        "new": curr_df[curr_df["身份证"].isin(new_ids)].to_dict("records"),
        "left": prev_df[prev_df["身份证"].isin(left_ids)].to_dict("records")
    }


def match_phones(persons, phone_dict, logger=None):
    """匹配电话号码"""
    def log(msg):
        if logger:
            logger(msg)
    
    matched = 0
    unmatched = []
    
    for person in persons:
        id_num = person["身份证"]
        if id_num in phone_dict:
            person["手机号"] = phone_dict[id_num]
            matched += 1
        else:
            person["手机号"] = ""
            unmatched.append(person)
    
    log(f"[INFO] 电话匹配: {matched} 条成功, {len(unmatched)} 条未匹配")
    
    if unmatched:
        log(f"[WARN] 以下人员未找到电话:")
        for p in unmatched:
            log(f"       - {p['姓名']} ({p['身份证']})")
    
    return unmatched


def generate_report(diff_result, phone_dict, output_file, logger=None):
    """生成Excel报告"""
    def log(msg):
        if logger:
            logger(msg)
    
    wb = Workbook()
    
    ws_new = wb.active
    ws_new.title = "花名册"
    ws_new.append(["*姓名", "*手机号", "*证件类型", "*证件号码"])
    for person in diff_result["new"]:
        phone = phone_dict.get(person["身份证"], "")
        ws_new.append([
            person["姓名"],
            phone,
            "居民身份证",
            person["身份证"]
        ])
    
    ws_left = wb.create_sheet("减少人员")
    ws_left.append(["姓名", "身份证"])
    for person in diff_result["left"]:
        ws_left.append([person["姓名"], person["身份证"]])
    
    ws_missing = wb.create_sheet("电话缺失")
    ws_missing.append(["姓名", "身份证", "缺失原因"])
    for person in diff_result["new"]:
        if person["身份证"] not in phone_dict:
            ws_missing.append([person["姓名"], person["身份证"], "未找到对应电话"])
    
    for sheet in wb.worksheets:
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length * 2 + 5, 80)
            sheet.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output_file)


def get_month_from_filename(filepath):
    """从文件名提取月份"""
    filename = Path(filepath).stem
    match = re.search(r"20\d{2}年?0?\d{1,2}月?", filename)
    if match:
        month_str = match.group()
        month_str = month_str.replace("年", "").replace("月", "")
        return month_str
    return datetime.now().strftime("%Y%m")


def run_comparison(prev_file, curr_file, phone_file, output_file, logger=None):
    """执行对比流程"""
    def log(msg):
        if logger:
            logger(msg)
    
    log("=" * 50)
    log(f"[INFO] 开始处理: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log("=" * 50)
    
    prev_df = read_salary_table(prev_file, logger)
    curr_df = read_salary_table(curr_file, logger)
    
    phone_dict = {}
    if phone_file and Path(phone_file).exists():
        phone_dict = read_phone_table(phone_file, logger)
    
    diff_result = compare_persons(prev_df, curr_df, logger)
    
    if diff_result["new"]:
        match_phones(diff_result["new"], phone_dict, logger)
    
    if not output_file:
        curr_month = get_month_from_filename(curr_file)
        output_file = f"工资人员差异报告_{curr_month}.xlsx"
    
    generate_report(diff_result, phone_dict, output_file, logger)
    
    log("=" * 50)
    log("[INFO] 处理完成!")
    log(f"[INFO] 新增 {len(diff_result['new'])} 人, 减少 {len(diff_result['left'])} 人")
    log("=" * 50)
    
    return diff_result


def create_gui():
    """创建GUI界面"""
    try:
        from tkinter import (
            Tk, Frame, Label, Entry, Button, Text, Scrollbar,
            StringVar, filedialog, messagebox, ttk
        )
        try:
            from tkinterdnd2 import Tk as DndTk
            HAS_DND = True
        except ImportError:
            HAS_DND = False
            DndTk = None
    except ImportError as e:
        print(f"GUI依赖库未安装: {e}")
        print("请运行: pip install tkinterdnd2")
        sys.exit(1)
    
    class PayrollComparatorGUI:
        def __init__(self, root):
            self.root = root
            self.root.title("工资表人员差异对比工具")
            self.root.geometry("900x700")
            self.root.resizable(True, True)
            self.has_dnd = HAS_DND
            
            self.log_text = None
            self.result_data = None
            
            self.setup_ui()
            self.load_saved_config()
            
            if not HAS_DND:
                self.log("[WARN] tkinterdnd2 未安装，拖拽功能不可用")
        
        def log(self, message):
            """添加日志"""
            if self.log_text:
                timestamp = datetime.now().strftime("%H:%M:%S")
                self.log_text.insert("end", f"{timestamp} {message}\n")
                self.log_text.see("end")
                self.root.update_idletasks()
        
        def load_saved_config(self):
            """加载保存的配置"""
            config = load_config()
            if "phone_file" in config and config["phone_file"]:
                self.phone_file.delete(0, "end")
                self.phone_file.insert(0, config["phone_file"])
                self.log("[INFO] 已加载保存的电话信息表路径")
        
        def save_config(self):
            """保存配置"""
            config = load_config()
            phone_path = self.phone_file.get()
            if phone_path:
                config["phone_file"] = phone_path
                save_config(config)
        
        def create_drop_target(self, parent, row, label_text):
            """创建支持拖拽的输入框"""
            label = Label(parent, text=label_text, anchor="e", width=15, font=("微软雅黑", 10))
            label.grid(row=row, column=0, sticky="e", padx=(10, 5), pady=5)
            
            entry = Entry(parent, font=("微软雅黑", 10), bd=2, relief="solid", 
                         highlightthickness=1, highlightbackground="#CCCCCC", width=30)
            entry.grid(row=row, column=1, sticky="ew", padx=5)
            
            btn = Button(parent, text="选择文件", font=("微软雅黑", 9), command=lambda: self.select_file(entry))
            btn.grid(row=row, column=2, padx=5)
            
            parent.grid_columnconfigure(1, weight=1)
            
            if self.has_dnd:
                try:
                    from tkinterdnd2 import DND_FILES
                    entry.drop_target_register(DND_FILES)
                    entry.dnd_bind('<<Drop>>', lambda e: self.handle_drop(e, entry))
                except Exception as e:
                    print(f"拖拽绑定失败: {e}")
            
            return entry
        
        def handle_drop(self, event, entry):
            """处理拖拽的文件"""
            try:
                from tkinterdnd2 import parse_dropfiles
                files = parse_dropfiles(event.data)
                if files:
                    filepath = files[0]
                    if filepath.endswith(('.xlsx', '.xls')):
                        entry.delete(0, "end")
                        entry.insert(0, filepath)
            except:
                path = event.data.strip()
                if path.startswith('{') and path.endswith('}'):
                    path = path[1:-1]
                if os.path.exists(path):
                    entry.delete(0, "end")
                    entry.insert(0, path)
        
        def select_file(self, entry):
            """选择文件"""
            filepath = filedialog.askopenfilename(
                title="选择文件",
                filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
            )
            if filepath:
                entry.delete(0, "end")
                entry.insert(0, filepath)
        
        def setup_ui(self):
            """设置UI"""
            main_frame = Frame(self.root, padx=20, pady=20)
            main_frame.grid(row=0, column=0, sticky="nsew")
            self.root.grid_rowconfigure(0, weight=1)
            self.root.grid_columnconfigure(0, weight=1)
            main_frame.grid_columnconfigure(0, weight=1)
            
            title_label = Label(
                main_frame,
                text="工资表人员差异对比工具",
                font=("微软雅黑", 16, "bold")
            )
            title_label.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 20))
            
            file_frame = Frame(main_frame, bd=2, relief="groove")
            file_frame.grid(row=1, column=0, columnspan=3, sticky="ew", pady=10, padx=0)
            
            file_frame.grid_columnconfigure(0, weight=0)
            file_frame.grid_columnconfigure(1, weight=1)
            file_frame.grid_columnconfigure(2, weight=0)
            
            hint_label = Label(file_frame, text="← 拖拽 Excel 文件到输入框", fg="#888888", font=("微软雅黑", 9))
            hint_label.grid(row=0, column=0, columnspan=3, sticky="w", padx=10, pady=(5, 0))
            
            self.prev_file = self.create_drop_target(file_frame, 1, "上月工资表:")
            self.curr_file = self.create_drop_target(file_frame, 2, "本月工资表:")
            self.phone_file = self.create_drop_target(file_frame, 3, "电话信息表:")
            
            btn_frame = Frame(main_frame)
            btn_frame.grid(row=2, column=0, columnspan=3, pady=15)
            
            self.run_btn = Button(
                btn_frame,
                text="开始对比",
                command=self.run_comparison,
                width=15,
                height=2,
                font=("微软雅黑", 11)
            )
            self.run_btn.grid(row=0, column=0, padx=10)
            
            self.export_btn = Button(
                btn_frame,
                text="打开结果文件夹",
                command=self.open_output_folder,
                width=15,
                height=2,
                font=("微软雅黑", 11),
                state="disabled"
            )
            self.export_btn.grid(row=0, column=1, padx=10)
            
            log_frame = Frame(main_frame)
            log_frame.grid(row=3, column=0, columnspan=3, sticky="nsew", pady=10)
            main_frame.grid_rowconfigure(3, weight=1)
            
            log_label = Label(log_frame, text="处理日志:", anchor="w", font=("微软雅黑", 10, "bold"))
            log_label.grid(row=0, column=0, sticky="w")
            
            self.log_text = Text(log_frame, height=12, font=("Consolas", 9))
            self.log_text.grid(row=1, column=0, sticky="nsew")
            log_frame.grid_rowconfigure(1, weight=1)
            log_frame.grid_columnconfigure(0, weight=1)
            
            scrollbar = Scrollbar(log_frame, command=self.log_text.yview)
            scrollbar.grid(row=1, column=1, sticky="ns")
            self.log_text.config(yscrollcommand=scrollbar.set)
            
            result_frame = Frame(main_frame)
            result_frame.grid(row=4, column=0, columnspan=3, sticky="nsew", pady=10)
            main_frame.grid_rowconfigure(4, weight=1)
            
            result_label = Label(result_frame, text="结果预览:", anchor="w", font=("微软雅黑", 10, "bold"))
            result_label.grid(row=0, column=0, sticky="w")
            
            self.result_tree = ttk.Treeview(result_frame, show="headings", height=8)
            self.result_tree.grid(row=1, column=0, sticky="nsew")
            result_frame.grid_rowconfigure(1, weight=1)
            result_frame.grid_columnconfigure(0, weight=1)
            
            scrollbar2 = Scrollbar(result_frame, command=self.result_tree.yview)
            scrollbar2.grid(row=1, column=1, sticky="ns")
            self.result_tree.config(yscrollcommand=scrollbar2.set)
            
            self.status_label = Label(
                main_frame,
                text="就绪",
                anchor="w",
                font=("微软雅黑", 9)
            )
            self.status_label.grid(row=5, column=0, columnspan=3, sticky="ew", pady=(10, 0))
        
        def validate_inputs(self):
            """验证输入"""
            errors = []
            
            if not self.prev_file.get():
                errors.append("请选择上月工资表")
            elif not os.path.exists(self.prev_file.get()):
                errors.append(f"上月工资表文件不存在: {self.prev_file.get()}")
            
            if not self.curr_file.get():
                errors.append("请选择本月工资表")
            elif not os.path.exists(self.curr_file.get()):
                errors.append(f"本月工资表文件不存在: {self.curr_file.get()}")
            
            if self.phone_file.get() and not os.path.exists(self.phone_file.get()):
                errors.append(f"电话信息表文件不存在: {self.phone_file.get()}")
            
            return errors
        
        def run_comparison(self):
            """执行对比"""
            errors = self.validate_inputs()
            if errors:
                for err in errors:
                    self.log(f"[ERROR] {err}")
                messagebox.showerror("输入错误", "\n".join(errors))
                return
            
            self.run_btn.config(state="disabled")
            self.log("-" * 50)
            
            try:
                self.result_data = run_comparison(
                    self.prev_file.get(),
                    self.curr_file.get(),
                    self.phone_file.get() if self.phone_file.get() else None,
                    None,
                    logger=self.log
                )
                
                self.save_config()
                self.update_result_preview()
                self.export_btn.config(state="normal")
                
                saved_path = self.auto_save_report()
                
                self.status_label.config(
                    text=f"完成: 新增 {len(self.result_data['new'])} 人, 减少 {len(self.result_data['left'])} 人"
                )
                
                messagebox.showinfo("完成", "对比完成，报告已保存")
            
            except Exception as e:
                self.log(f"[ERROR] 处理失败: {str(e)}")
                messagebox.showerror("错误", f"处理失败:\n{str(e)}")
            
            finally:
                self.run_btn.config(state="normal")
        
        def auto_save_report(self):
            """自动保存报告"""
            curr_path = self.curr_file.get()
            if not curr_path:
                return
            
            base_dir = os.path.dirname(curr_path)
            output_dir = os.path.join(base_dir, "difference")
            os.makedirs(output_dir, exist_ok=True)
            
            curr_filename = os.path.basename(curr_path)
            curr_name = os.path.splitext(curr_filename)[0]
            
            filename = f"{curr_name}_工资人员差异报告_{datetime.now().strftime('%Y%m%d')}.xlsx"
            filepath = os.path.join(output_dir, filename)
            
            counter = 1
            while os.path.exists(filepath):
                filename = f"{curr_name}_工资人员差异报告_{datetime.now().strftime('%Y%m%d')}_{counter}.xlsx"
                filepath = os.path.join(output_dir, filename)
                counter += 1
            
            phone_dict = {}
            if self.phone_file.get() and os.path.exists(self.phone_file.get()):
                phone_dict = read_phone_table(self.phone_file.get())
                phone_dict_clean = {name: str(phone).rstrip('.0') for name, phone in phone_dict.items() if phone}
                phone_dict = phone_dict_clean
            
            generate_report(self.result_data, phone_dict, filepath, logger=self.log)
            self.log(f"[INFO] 报告已保存: {os.path.abspath(filepath)}")
            
            return filepath
        
        def update_result_preview(self):
            """更新结果预览"""
            for item in self.result_tree.get_children():
                self.result_tree.delete(item)
            
            self.result_tree["columns"] = ["姓名", "身份证", "手机号", "状态"]
            self.result_tree.heading("姓名", text="姓名")
            self.result_tree.heading("身份证", text="身份证")
            self.result_tree.heading("手机号", text="手机号")
            self.result_tree.heading("状态", text="状态")
            
            self.result_tree.column("姓名", width=100)
            self.result_tree.column("身份证", width=180)
            self.result_tree.column("手机号", width=120)
            self.result_tree.column("状态", width=80)
            
            if self.result_data:
                for person in self.result_data["new"]:
                    phone = person.get("手机号", "")
                    status = "电话缺失" if not phone else "新增"
                    self.result_tree.insert("", "end", values=[
                        person["姓名"], person["身份证"], phone, status
                    ])
                
                for person in self.result_data["left"]:
                    self.result_tree.insert("", "end", values=[
                        person["姓名"], person["身份证"], "-", "减少"
                    ])
        
        def export_report(self):
            """导出报告"""
            if not self.result_data:
                messagebox.showwarning("警告", "没有可导出的数据")
                return
            
            curr_path = self.curr_file.get()
            if not curr_path:
                messagebox.showwarning("警告", "请先选择本月工资表")
                return
            
            base_dir = os.path.dirname(curr_path)
            output_dir = os.path.join(base_dir, "difference")
            os.makedirs(output_dir, exist_ok=True)
            
            filename = f"工资人员差异报告_{datetime.now().strftime('%Y%m%d')}.xlsx"
            filepath = os.path.join(output_dir, filename)
            
            counter = 1
            while os.path.exists(filepath):
                filename = f"工资人员差异报告_{datetime.now().strftime('%Y%m%d')}_{counter}.xlsx"
                filepath = os.path.join(output_dir, filename)
                counter += 1
            
            try:
                phone_dict = {}
                if self.phone_file.get() and os.path.exists(self.phone_file.get()):
                    phone_dict = read_phone_table(self.phone_file.get())
                
                phone_dict_clean = {}
                for name, phone in phone_dict.items():
                    if phone:
                        phone_str = str(phone)
                        if '.' in phone_str and phone_str.endswith('.0'):
                            phone_str = phone_str.replace('.0', '')
                        phone_dict_clean[name] = phone_str
                phone_dict = phone_dict_clean
                
                generate_report(self.result_data, phone_dict, filepath, logger=self.log)
                self.log(f"[INFO] 报告已保存: {os.path.abspath(filepath)}")
                messagebox.showinfo("成功", f"报告已保存:\n{filepath}")
                
                if sys.platform == "win32":
                    os.startfile(output_dir)
                elif sys.platform == "darwin":
                    os.system(f"open '{output_dir}'")
                else:
                    os.system(f"xdg-open '{output_dir}'")
            
            except Exception as e:
                self.log(f"[ERROR] 导出失败: {str(e)}")
                messagebox.showerror("错误", f"导出失败:\n{str(e)}")
        
        def open_output_folder(self):
            """打开结果文件夹"""
            curr_path = self.curr_file.get()
            if not curr_path:
                messagebox.showwarning("警告", "请先选择本月工资表")
                return
            
            base_dir = os.path.dirname(curr_path)
            output_dir = os.path.join(base_dir, "difference")
            
            if not os.path.exists(output_dir):
                messagebox.showwarning("警告", "结果文件夹不存在，请先运行对比")
                return
            
            if sys.platform == "win32":
                os.startfile(output_dir)
            elif sys.platform == "darwin":
                os.system(f"open '{output_dir}'")
            else:
                os.system(f"xdg-open '{output_dir}'")
        
        def on_closing(self):
            """关闭窗口时保存配置"""
            self.save_config()
            self.root.destroy()
    
    root = DndTk() if HAS_DND else Tk()
    app = PayrollComparatorGUI(root)
    root.protocol("WM_DELETE_WINDOW", app.on_closing)
    root.mainloop()


def main():
    parser = argparse.ArgumentParser(
        description="工资表人员差异对比工具",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  # GUI模式
  python PayrollComparator.py
  
  # 命令行模式
  python PayrollComparator.py -p 上月.xlsx -c 本月.xlsx -t 电话.xlsx -o 输出.xlsx
        """
    )
    
    parser.add_argument("--gui", action="store_true", help="启动GUI模式")
    
    parser.add_argument("-p", "--prev", dest="prev_file", 
                       help="上月工资表文件路径")
    parser.add_argument("-c", "--curr", dest="curr_file",
                       help="本月工资表文件路径")
    parser.add_argument("-t", "--phone", dest="phone_file",
                       help="电话信息表文件路径")
    parser.add_argument("-o", "--output", dest="output_file",
                       help="输出报告文件路径")
    
    args = parser.parse_args()
    
    if args.gui or not any([args.prev_file, args.curr_file]):
        create_gui()
    else:
        if not args.prev_file or not args.curr_file:
            parser.print_help()
            print("\n错误: 命令行模式需要指定 -p 和 -c 参数")
            sys.exit(1)
        
        def cli_logger(msg):
            print(msg)
        
        run_comparison(
            args.prev_file,
            args.curr_file,
            args.phone_file,
            args.output_file,
            logger=cli_logger
        )


if __name__ == "__main__":
    main()
